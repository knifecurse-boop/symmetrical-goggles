"""把試卷排版輸出成 Markdown、HTML（可直接列印）與 Word（.docx）。

題號跨大題連續編號；克漏字與文意選填文章中的空格也會改成對應的題號。
每種格式都有「學生卷」與「教師卷（含答案與解析）」兩個版本。
"""

from __future__ import annotations

import html
import re
from dataclasses import dataclass
from pathlib import Path

from .models import Exam, ExamSection, Question
from .question_types import LETTERS, get_type

_BLANK_RE = re.compile(r"__\((\d+)\)__")


def _fmt_points(p: float) -> str:
    return str(int(p)) if float(p).is_integer() else str(p)


@dataclass
class NumberedSection:
    section: ExamSection
    start: int  # 本大題第一題的題號

    @property
    def layout(self) -> str:
        return get_type(self.section.type).layout

    @property
    def questions(self) -> list[tuple[int, Question]]:
        return [(self.start + i, q) for i, q in enumerate(self.section.content.questions)]

    @property
    def header(self) -> str:
        n = len(self.section.content.questions)
        p = self.section.points
        return f"{self.section.title}（每題 {_fmt_points(p)} 分，共 {_fmt_points(p * n)} 分）"

    def passage(self, blank: str = "__{n}__") -> str:
        text = self.section.content.passage or ""
        return _BLANK_RE.sub(lambda m: blank.format(n=self.start + int(m.group(1)) - 1), text)

    def stem(self, q: Question) -> str:
        # 克漏字/文意選填的題幹只是空格編號，在試卷上不另外印出
        if self.section.type in ("cloze", "word_bank"):
            return ""
        return q.stem


def number_sections(exam: Exam) -> list[NumberedSection]:
    result, n = [], 1
    for s in exam.sections:
        result.append(NumberedSection(s, n))
        n += len(s.content.questions)
    return result


# ---------------------------------------------------------------------------
# Markdown
# ---------------------------------------------------------------------------


def to_markdown(exam: Exam, answers: bool = False) -> str:
    out = [f"# {exam.title}" + ("（教師卷）" if answers else "")]
    if exam.subtitle:
        out.append(f"範圍：{exam.subtitle}　總分：{_fmt_points(exam.total_points)} 分")
    if not answers:
        out.append("班級：________　座號：____　姓名：____________　得分：______")
    for ns in number_sections(exam):
        out.append(f"\n## {ns.header}\n\n{ns.section.directions}")
        content = ns.section.content
        if content.word_bank:
            out.append("\n" + "　".join(f"({LETTERS[i]}) {w}" for i, w in enumerate(content.word_bank)))
        if content.passage:
            out.append("\n" + ns.passage("___{n}___"))
        out.append("")
        for num, q in ns.questions:
            stem = ns.stem(q)
            line = f"{num}. {stem}" if stem else f"{num}."
            if ns.layout == "open" and not answers:
                line += "\n\n   ______________________________________________"
            out.append(line)
            if q.options:
                out.append("   " + "　".join(f"({LETTERS[i]}) {o}" for i, o in enumerate(q.options)))
            if answers:
                out.append(f"   **答案：{q.answer}**　{q.explanation}（{q.source}）")
            out.append("")
    if not answers:
        return "\n".join(out).rstrip() + "\n"
    out.append("\n## 答案速查\n")
    out.append(" ".join(f"{num}.{q.answer}" for ns in number_sections(exam) if ns.layout != "open" for num, q in ns.questions))
    return "\n".join(out).rstrip() + "\n"


# ---------------------------------------------------------------------------
# HTML（可列印）
# ---------------------------------------------------------------------------

HTML_STYLE = """
:root { --ink:#1a1a1a; --muted:#555; --accent:#1f5fa8; --line:#ccc; --key:#b3261e; --bg:#fff; }
body { background:var(--bg); color:var(--ink); font-family:"Times New Roman","Noto Serif TC","PMingLiU",serif;
       max-width:820px; margin:24px auto; padding:0 16px; line-height:1.7; font-size:16px; }
h1 { text-align:center; font-size:24px; margin-bottom:4px; }
.meta { text-align:center; color:var(--muted); font-size:14px; }
.student { display:flex; justify-content:space-between; border-bottom:1px solid var(--line); padding:8px 0 12px; margin-bottom:8px; }
h2 { font-size:18px; margin:24px 0 4px; border-left:4px solid var(--accent); padding-left:8px; }
.directions { color:var(--muted); font-size:14px; margin:0 0 8px; }
.passage { text-align:justify; text-indent:2em; margin:8px 0 12px; white-space:pre-line; }
.blank { display:inline-block; min-width:3.2em; border-bottom:1px solid var(--ink); text-align:center; text-indent:0; font-weight:bold; }
.bank { border:1px solid var(--ink); padding:6px 10px; display:flex; flex-wrap:wrap; gap:4px 18px; margin:8px 0; }
.q { margin:6px 0 10px; break-inside:avoid; }
.opts { display:flex; flex-wrap:wrap; gap:2px 22px; padding-left:2em; }
.writeline { border-bottom:1px solid var(--ink); height:2em; margin:4px 0 4px 2em; }
.ans { color:var(--key); font-size:14px; padding-left:2em; }
.ans b { font-size:15px; }
.quick { font-family:monospace; font-size:14px; line-height:1.9; }
@media print { body { margin:0 auto; } h2 { break-after:avoid; } }
"""


def _e(s: str) -> str:
    return html.escape(s, quote=False)


def to_html(exam: Exam, answers: bool = False) -> str:
    title = exam.title + ("（教師卷）" if answers else "")
    parts = [
        "<!doctype html><html lang='zh-Hant'><head><meta charset='utf-8'>",
        "<meta name='viewport' content='width=device-width, initial-scale=1'>",
        f"<title>{_e(title)}</title><style>{HTML_STYLE}</style></head><body>",
        f"<h1>{_e(title)}</h1>",
    ]
    if exam.subtitle:
        parts.append(f"<p class='meta'>範圍：{_e(exam.subtitle)}　總分：{_fmt_points(exam.total_points)} 分</p>")
    if not answers:
        parts.append(
            "<div class='student'><span>班級：________</span><span>座號：____</span>"
            "<span>姓名：____________</span><span>得分：______</span></div>"
        )
    sections = number_sections(exam)
    for ns in sections:
        content = ns.section.content
        parts.append(f"<h2>{_e(ns.header)}</h2><p class='directions'>{_e(ns.section.directions)}</p>")
        if content.word_bank:
            parts.append(
                "<div class='bank'>"
                + "".join(f"<span>({LETTERS[i]}) {_e(w)}</span>" for i, w in enumerate(content.word_bank))
                + "</div>"
            )
        if content.passage:
            passage = _BLANK_RE.sub(
                lambda m: f"\x00{ns.start + int(m.group(1)) - 1}\x01", content.passage
            )
            passage = re.sub(r"\x00(\d+)\x01", r"<span class='blank'>\1</span>", _e(passage))
            parts.append(f"<div class='passage'>{passage}</div>")
        for num, q in ns.questions:
            stem = ns.stem(q)
            parts.append(f"<div class='q'><div>{num}. {_e(stem)}</div>")
            if q.options:
                parts.append(
                    "<div class='opts'>"
                    + "".join(f"<span>({LETTERS[i]}) {_e(o)}</span>" for i, o in enumerate(q.options))
                    + "</div>"
                )
            if ns.layout == "open" and not answers:
                parts.append("<div class='writeline'></div>")
            if answers:
                parts.append(
                    f"<div class='ans'><b>答案：{_e(q.answer)}</b>　{_e(q.explanation)}"
                    f"<br><small>命題依據：{_e(q.source)}</small></div>"
                )
            parts.append("</div>")
    if answers:
        quick = " ".join(
            f"{num}.{_e(q.answer)}" for ns in sections if ns.layout != "open" for num, q in ns.questions
        )
        parts.append(f"<h2>答案速查</h2><p class='quick'>{quick}</p>")
    parts.append("</body></html>")
    return "\n".join(parts)


# ---------------------------------------------------------------------------
# Word (.docx)
# ---------------------------------------------------------------------------


def to_docx(exam: Exam, path: Path | str, answers: bool = False) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "新細明體")

    def para(text: str = "", bold: bool = False, size: int | None = None, center: bool = False, color=None):
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        if color:
            run.font.color.rgb = color
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        return p

    para(exam.title + ("（教師卷）" if answers else ""), bold=True, size=18, center=True)
    if exam.subtitle:
        para(f"範圍：{exam.subtitle}　總分：{_fmt_points(exam.total_points)} 分", center=True, size=10)
    if not answers:
        para("班級：________　座號：____　姓名：____________　得分：______")

    red = RGBColor(0xB3, 0x26, 0x1E)
    for ns in number_sections(exam):
        content = ns.section.content
        para()
        para(ns.header, bold=True, size=13)
        para(ns.section.directions, size=10)
        if content.word_bank:
            para("　".join(f"({LETTERS[i]}) {w}" for i, w in enumerate(content.word_bank)))
        if content.passage:
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(24)
            pos = 0
            for m in _BLANK_RE.finditer(content.passage):
                p.add_run(content.passage[pos : m.start()])
                blank = p.add_run(f"  {ns.start + int(m.group(1)) - 1}  ")
                blank.underline = True
                blank.bold = True
                pos = m.end()
            p.add_run(content.passage[pos:])
        for num, q in ns.questions:
            stem = ns.stem(q)
            para(f"{num}. {stem}" if stem else f"{num}.")
            if q.options:
                para("    " + "　".join(f"({LETTERS[i]}) {o}" for i, o in enumerate(q.options)))
            if ns.layout == "open" and not answers:
                para("    " + "_" * 60)
            if answers:
                para(f"    答案：{q.answer}　{q.explanation}（{q.source}）", size=10, color=red)
    doc.save(str(path))


def write_all(exam: Exam, out_dir: Path | str, stem: str = "exam") -> list[Path]:
    """輸出學生卷與教師卷的 HTML、Markdown、Word 以及原始 JSON，回傳寫出的檔案路徑。"""
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    written = []
    for answers, suffix in ((False, ""), (True, "_answers")):
        for ext, render in (("html", to_html), ("md", to_markdown)):
            p = out / f"{stem}{suffix}.{ext}"
            p.write_text(render(exam, answers=answers), encoding="utf-8")
            written.append(p)
        p = out / f"{stem}{suffix}.docx"
        to_docx(exam, p, answers=answers)
        written.append(p)
    p = out / f"{stem}.json"
    p.write_text(exam.model_dump_json(indent=2), encoding="utf-8")
    written.append(p)
    return written
